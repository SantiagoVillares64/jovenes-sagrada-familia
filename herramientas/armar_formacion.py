"""Arma js/formacion-datos.js a partir de banco_formacion.py y deja un Excel para revisar las preguntas.

Uso:  python herramientas/armar_formacion.py
"""
import json
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
WEB = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
from banco_formacion import CURSOS, PROXIMOS  # noqa: E402


def controlar():
    ids = [c['id'] for c in CURSOS] + [c['id'] for c in PROXIMOS]
    if len(ids) != len(set(ids)):
        sys.exit('Hay cursos con el mismo id')
    for c in CURSOS:
        if len(c['preguntas']) < 10:
            sys.exit(f'{c["titulo"]}: tiene {len(c["preguntas"])} preguntas (mínimo 10)')
        vistas = set()
        for q in c['preguntas']:
            opciones = [q['ok']] + q['otras']
            if len(q['otras']) != 3 or len(set(opciones)) != 4:
                sys.exit(f'{c["titulo"]}: «{q["p"]}» necesita 1 correcta y 3 incorrectas distintas')
            if not q.get('ref'):
                sys.exit(f'{c["titulo"]}: «{q["p"]}» no tiene referencia')
            if q['p'] in vistas:
                sys.exit(f'{c["titulo"]}: pregunta repetida «{q["p"]}»')
            vistas.add(q['p'])


def excel():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return None
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for c in CURSOS:
        ws = wb.create_sheet(c['titulo'][:30].replace('?', '').replace('¿', ''))
        ws.append(['#', 'Pregunta', 'Respuesta correcta', 'Incorrecta 1', 'Incorrecta 2', 'Incorrecta 3', 'Fuente', '¿OK? / Comentario'])
        for i, q in enumerate(c['preguntas'], 1):
            ws.append([i, q['p'], q['ok']] + q['otras'] + [q['ref'], ''])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col, w in zip('ABCDEFGH', [4, 48, 42, 28, 28, 28, 16, 28]):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'B2'
    ruta = os.path.join(os.path.dirname(WEB), 'Para revisar', 'Formación - preguntas para revisar.xlsx')
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    wb.save(ruta)
    return ruta


def main():
    controlar()
    datos = {'cursos': CURSOS, 'proximos': PROXIMOS}
    js = ('/* Cursos de la Academia Frassati.\n'
          '   NO editar a mano: se genera con  python herramientas/armar_formacion.py */\n'
          'window.FORMACION = ' + json.dumps(datos, ensure_ascii=False, separators=(',', ':')) + ';\n')
    open(os.path.join(WEB, 'js', 'formacion-datos.js'), 'w', encoding='utf-8', newline='\n').write(js)
    ruta = excel()
    print(f'Listo: {len(CURSOS)} cursos (' + ', '.join(f'{c["titulo"]}: {len(c["preguntas"])}' for c in CURSOS) +
          f') + {len(PROXIMOS)} próximamente · {round(len(js.encode()) / 1024)} KB')
    if ruta:
        print('  Excel para revisar:', ruta)


if __name__ == '__main__':
    main()
